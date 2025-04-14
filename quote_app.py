import streamlit as st
import pandas as pd
import openpyxl
from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from PIL import Image
import io
from fpdf import FPDF

# Load CPI table
cpi_df = pd.read_excel("Refined_CPI_Lookup_Table.xlsx", index_col=0)

# Programming logic
prog_rates = {
    "Simple": (350, 50),
    "Standard": (500, 75),
    "Complex": (800, 100)
}

def calculate_quote(loi, ir, n, prog_type, dp_type):
    cpi = round(round(cpi_df.loc[loi, ir] / 0.05) * 0.05, 2)
    sample_cost = cpi * n
    base, per_min = prog_rates[prog_type]
    prog_cost = base + per_min * loi
    dp_cost = 180 if dp_type == "SPSS" else loi * 50
    total = sample_cost + prog_cost + dp_cost
    return cpi, sample_cost, prog_cost, dp_cost, total

def generate_pdf(quote_df, audience_df, client, contact, company, ref, notes, project_name):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.set_font("Arial", "B", 16)
    pdf.cell(0, 10, "PureSpectrum Multi-Option Quote", ln=True, align="C")
    pdf.ln(5)

    pdf.set_font("Arial", size=12)
    for label, val in [
        ("Client", client),
        ("Contact", contact),
        ("Company", company),
        ("Reference", ref),
        ("Project Name", project_name)
    ]:
        pdf.cell(50, 10, f"{label}:", ln=0)
        pdf.cell(0, 10, str(val), ln=1)

    if notes:
        pdf.ln(5)
        pdf.set_font("Arial", "B", 12)
        pdf.cell(0, 10, "Notes:", ln=True)
        pdf.set_font("Arial", size=12)
        pdf.multi_cell(0, 8, notes)

    if not audience_df.empty:
        pdf.ln(5)
        pdf.set_font("Arial", "B", 12)
        pdf.cell(0, 10, "Audience:", ln=True)
        pdf.set_font("Arial", size=12)
        for _, row in audience_df.iterrows():
            pdf.multi_cell(0, 8, f"- {row['Audience Description']}")

    pdf.ln(5)
    pdf.set_font("Arial", "B", 12)
    pdf.cell(0, 10, "Quote Options:", ln=True)

    pdf.set_font("Arial", "B", 10)
    headers = list(quote_df.columns)
    col_widths = [pdf.get_string_width(h) + 6 for h in headers]
    for i, h in enumerate(headers):
        pdf.cell(col_widths[i], 8, h, border=1)
    pdf.ln()
    pdf.set_font("Arial", size=10)
    for _, row in quote_df.iterrows():
        for i, h in enumerate(headers):
            pdf.cell(col_widths[i], 8, str(row[h]), border=1)
        pdf.ln()

    pdf.ln(5)
    pdf.set_font("Arial", "I", 10)
    pdf.cell(0, 10, "Prepared by PureSpectrum — All quotes valid for 30 days", ln=True, align="C")

    buf = io.BytesIO()
    pdf.output(buf)
    buf.seek(0)
    return buf

st.title("PureSpectrum Multi-Option Quote Builder")

with st.form("project_form"):
    client = st.text_input("Client Name")
    contact = st.text_input("Contact Person")
    company = st.text_input("Company Name")
    ref = st.text_input("Reference Number")
    prog = st.selectbox("Programming Type", ["Simple", "Standard", "Complex"])
    dp = st.selectbox("DP Type", ["SPSS", "Tables"])
    notes = st.text_area("Quoting Notes (optional)")

    st.markdown("### Target Audience")
    audience_rows = st.experimental_data_editor(pd.DataFrame({"Audience Description": [""]}), num_rows="dynamic")

    st.markdown("### Quote Options")
    option_rows = st.experimental_data_editor(
        pd.DataFrame({"LOI": [10], "IR": [50], "n": [1000]}),
        num_rows="dynamic")

    submit = st.form_submit_button("Generate Quote")

if submit:
    summary = []
    for _, row in option_rows.iterrows():
        loi, ir, n = int(row["LOI"]), int(row["IR"]), int(row["n"])
        cpi, sample, prog_cost, dp_cost, total = calculate_quote(loi, ir, n, prog, dp)
        summary.append({
            "LOI (min)": loi,
            "IR (%)": ir,
            "n": n,
            "CPI": f"${cpi:.2f}",
            "Sample Cost": f"${sample:,.2f}",
            "Programming Cost": f"${prog_cost:,.2f}",
            "DP Cost": f"${dp_cost:,.2f}",
            "Total": f"${total:,.2f}"
        })

    quote_df = pd.DataFrame(summary)
    st.subheader("Quote Summary")
    st.dataframe(quote_df)

    unique_ns = sorted(set(option_rows['n']))
    unique_lois = sorted(set(option_rows['LOI']))
    unique_irs = sorted(set(option_rows['IR']))
    project_name = f"n={','.join(map(str, unique_ns))} | LOI={','.join(map(str, unique_lois))} | IR={','.join(f'{i}%' for i in unique_irs)}"
    st.write(f"**Generated Project Name:** {project_name}")

    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        quote_df.to_excel(writer, sheet_name="Quote", index=False)
        audience_rows.to_excel(writer, sheet_name="Audience", index=False)
        workbook = writer.book
        ws = writer.sheets["Quote"]
        ws.cell(row=1, column=quote_df.shape[1]+2, value="Client:").offset(0, 1).value = client
        ws.cell(row=2, column=quote_df.shape[1]+2, value="Contact:").offset(0, 1).value = contact
        ws.cell(row=3, column=quote_df.shape[1]+2, value="Company:").offset(0, 1).value = company
        ws.cell(row=4, column=quote_df.shape[1]+2, value="Ref:").offset(0, 1).value = ref
        ws.cell(row=5, column=quote_df.shape[1]+2, value="Notes:").offset(0, 1).value = notes
        ws.cell(row=6, column=quote_df.shape[1]+2, value="Project Name:").offset(0, 1).value = project_name
    output.seek(0)

    pdf_data = generate_pdf(quote_df, audience_rows, client, contact, company, ref, notes, project_name)

    st.download_button("Download Excel Quote", data=output, file_name=f"Quote_{ref or 'project'}.xlsx", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    st.download_button("Download PDF Quote", data=pdf_data, file_name=f"Quote_{ref or 'project'}.pdf", mime="application/pdf")
